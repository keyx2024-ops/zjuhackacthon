"""
文件解析模块

支持多种文件格式：PDF、DOCX、Markdown、TXT
自动识别章节结构，提取文本内容和元数据
"""
import logging
import re
import uuid
from pathlib import Path
from typing import List, Optional, Tuple

import fitz
from docx import Document as DocxDocument

from models.textbook import Chapter, FileFormat, Textbook, TextbookStatus


logger = logging.getLogger(__name__)


PDF_CHAPTER_PATTERNS = [
    r"^第([一二三四五六七八九十百千]+|\d+)[章部篇]\s*[一-鿿]",
    r"^Chapter\s+(\d+)",
]

TXT_CHAPTER_PATTERNS = PDF_CHAPTER_PATTERNS + [
    r"^第([一二三四五六七八九十百千]+|\d+)节",
    r"^Section\s+(\d+)",
    r"^[#]+\s+",
]

NON_CONTENT_TITLE_RE = re.compile(
    r"^(封面|书名|版权|编委名单|主编简介|副主编简介|编者|"
    r"教材修订说明|新形态教材使用说明|致谢|前言|序言?|"
    r"目\s*录|参考文献|附录|索引|版次|印次)"
)
MAX_CHAPTERS_PER_BOOK = 60


class FileParser:
    """文件解析器"""

    def __init__(self):
        self.parsers = {
            FileFormat.PDF: self._parse_pdf,
            FileFormat.DOCX: self._parse_docx,
            FileFormat.MD: self._parse_markdown,
            FileFormat.TXT: self._parse_txt,
        }

    def parse(self, file_path: str, textbook_name: Optional[str] = None) -> Textbook:
        """
        解析文件并返回结构化教材数据

        Args:
            file_path: 文件路径
            textbook_name: 教材名称（可选，默认使用文件名）

        Returns:
            Textbook: 结构化教材数据

        Raises:
            ValueError: 文件格式不支持
            FileNotFoundError: 文件不存在
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_format = self._detect_format(path)
        if file_format is None:
            raise ValueError(f"Unsupported file format: {path.suffix}")

        if textbook_name is None:
            textbook_name = path.stem

        logger.info(f"Parsing {file_format.value} file: {path.name}")

        parser = self.parsers[file_format]
        chapters, total_pages, metadata = parser(path)

        total_words = sum(c.word_count for c in chapters)

        textbook = Textbook(
            textbook_id=str(uuid.uuid4()),
            name=textbook_name,
            file_format=file_format,
            file_size=path.stat().st_size,
            file_path=str(path.absolute()),
            status=TextbookStatus.PARSED,
            total_pages=total_pages,
            total_words=total_words,
            chapters=chapters,
            metadata=metadata,
        )

        logger.info(
            f"Parsed {textbook_name}: {len(chapters)} chapters, "
            f"{total_pages} pages, {total_words} words"
        )

        return textbook

    def _detect_format(self, path: Path) -> Optional[FileFormat]:
        """检测文件格式"""
        suffix = path.suffix.lower().lstrip(".")
        format_map = {
            "pdf": FileFormat.PDF,
            "docx": FileFormat.DOCX,
            "md": FileFormat.MD,
            "markdown": FileFormat.MD,
            "txt": FileFormat.TXT,
        }
        return format_map.get(suffix)

    def _parse_pdf(self, path: Path) -> Tuple[List[Chapter], int, dict]:
        """解析 PDF 文件"""
        doc = fitz.open(str(path))
        total_pages = doc.page_count

        all_text = []
        for page_num in range(total_pages):
            page = doc[page_num]
            text = page.get_text()
            text = self._clean_pdf_text(text)
            all_text.append((page_num + 1, text))

        toc = doc.get_toc()
        metadata = {
            "title": doc.metadata.get("title", ""),
            "author": doc.metadata.get("author", ""),
            "subject": doc.metadata.get("subject", ""),
            "has_toc": len(toc) > 0,
        }

        if toc:
            chapters = self._parse_chapters_with_toc(toc, all_text)
        else:
            chapters = self._parse_chapters_with_pattern(all_text, PDF_CHAPTER_PATTERNS)

        chapters = self._filter_and_cap(chapters)

        doc.close()
        return chapters, total_pages, metadata

    def _clean_pdf_text(self, text: str) -> str:
        """清理 PDF 文本（去除页眉页脚等）"""
        lines = text.split("\n")
        cleaned_lines = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if re.match(r"^[\d\s\-·•—]+$", line):
                continue
            if len(line) < 3 and line.isdigit():
                continue
            cleaned_lines.append(line)
        return "\n".join(cleaned_lines)

    def _parse_chapters_with_toc(
        self, toc: list, all_text: List[Tuple[int, str]]
    ) -> List[Chapter]:
        """使用目录解析章节"""
        chapters = []
        text_by_page = {p: t for p, t in all_text}

        top_level_toc = [item for item in toc if item[0] == 1]
        if not top_level_toc:
            top_level_toc = toc

        for i, item in enumerate(top_level_toc):
            level, title, page_start = item
            page_end = (
                top_level_toc[i + 1][2] - 1
                if i + 1 < len(top_level_toc)
                else max(p for p, _ in all_text)
            )

            content_parts = []
            for page_num in range(page_start, page_end + 1):
                if page_num in text_by_page:
                    content_parts.append(text_by_page[page_num])
            content = "\n".join(content_parts)

            chapter = Chapter(
                chapter_id=str(uuid.uuid4()),
                chapter_number=str(i + 1),
                title=title.strip(),
                content=content,
                page_start=page_start,
                page_end=page_end,
                word_count=len(content),
            )
            chapters.append(chapter)

        return chapters

    def _parse_chapters_with_pattern(
        self, all_text: List[Tuple[int, str]], patterns: list
    ) -> List[Chapter]:
        """使用正则表达式解析章节"""
        full_text = "\n".join(text for _, text in all_text)
        page_offsets = self._compute_page_offsets(all_text)

        chapter_marks = []
        for pattern in patterns:
            for match in re.finditer(pattern, full_text, re.MULTILINE):
                chapter_marks.append((match.start(), match.group(0)))

        chapter_marks.sort()

        if not chapter_marks:
            return [
                Chapter(
                    chapter_id=str(uuid.uuid4()),
                    chapter_number="1",
                    title="全文",
                    content=full_text,
                    page_start=1,
                    page_end=max(p for p, _ in all_text),
                    word_count=len(full_text),
                )
            ]

        chapters = []
        for i, (start, title) in enumerate(chapter_marks):
            end = chapter_marks[i + 1][0] if i + 1 < len(chapter_marks) else len(full_text)
            content = full_text[start:end].strip()
            page_start = self._find_page_for_offset(start, page_offsets)
            page_end = self._find_page_for_offset(end - 1, page_offsets)

            chapter = Chapter(
                chapter_id=str(uuid.uuid4()),
                chapter_number=str(i + 1),
                title=title.strip(),
                content=content,
                page_start=page_start,
                page_end=page_end,
                word_count=len(content),
            )
            chapters.append(chapter)

        return chapters

    def _compute_page_offsets(
        self, all_text: List[Tuple[int, str]]
    ) -> List[Tuple[int, int]]:
        """计算每页文本的起始偏移量"""
        offsets = []
        current_offset = 0
        for page_num, text in all_text:
            offsets.append((current_offset, page_num))
            current_offset += len(text) + 1
        return offsets

    def _find_page_for_offset(
        self, offset: int, page_offsets: List[Tuple[int, int]]
    ) -> int:
        """根据偏移量找到对应的页码"""
        for i, (page_offset, page_num) in enumerate(page_offsets):
            if i + 1 < len(page_offsets):
                next_offset = page_offsets[i + 1][0]
                if page_offset <= offset < next_offset:
                    return page_num
            else:
                return page_num
        return 1

    def _parse_docx(self, path: Path) -> Tuple[List[Chapter], int, dict]:
        """解析 DOCX 文件"""
        doc = DocxDocument(str(path))
        chapters = []
        current_chapter = None
        current_content = []
        chapter_number = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            is_heading = para.style.name.startswith("Heading")
            if is_heading and (para.style.name == "Heading 1" or para.style.name == "Heading 2"):
                if current_chapter is not None:
                    content = "\n".join(current_content)
                    current_chapter.content = content
                    current_chapter.word_count = len(content)
                    chapters.append(current_chapter)

                chapter_number += 1
                current_chapter = Chapter(
                    chapter_id=str(uuid.uuid4()),
                    chapter_number=str(chapter_number),
                    title=text,
                    content="",
                    word_count=0,
                )
                current_content = []
            else:
                current_content.append(text)

        if current_chapter is not None:
            content_full = "\n".join(current_content)
            current_chapter.content = content_full
            current_chapter.word_count = len(content_full)
            chapters.append(current_chapter)

        if not chapters:
            full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            chapters = [
                Chapter(
                    chapter_id=str(uuid.uuid4()),
                    chapter_number="1",
                    title="全文",
                    content=full_text,
                    word_count=len(full_text),
                )
            ]

        chapters = self._filter_and_cap(chapters)
        metadata = {
            "title": doc.core_properties.title or "",
            "author": doc.core_properties.author or "",
        }

        return chapters, len(chapters), metadata

    def _parse_markdown(self, path: Path) -> Tuple[List[Chapter], int, dict]:
        """解析 Markdown 文件"""
        content = path.read_text(encoding="utf-8")
        chapters = []
        current_chapter = None
        current_content = []
        chapter_number = 0

        for line in content.split("\n"):
            heading_match = re.match(r"^(#{1,2})\s+(.+)$", line)
            if heading_match:
                if current_chapter is not None:
                    chapter_content = "\n".join(current_content)
                    current_chapter.content = chapter_content
                    current_chapter.word_count = len(chapter_content)
                    chapters.append(current_chapter)

                chapter_number += 1
                current_chapter = Chapter(
                    chapter_id=str(uuid.uuid4()),
                    chapter_number=str(chapter_number),
                    title=heading_match.group(2).strip(),
                    content="",
                    word_count=0,
                )
                current_content = []
            else:
                current_content.append(line)

        if current_chapter is not None:
            chapter_content = "\n".join(current_content)
            current_chapter.content = chapter_content
            current_chapter.word_count = len(chapter_content)
            chapters.append(current_chapter)

        if not chapters:
            chapters = [
                Chapter(
                    chapter_id=str(uuid.uuid4()),
                    chapter_number="1",
                    title="全文",
                    content=content,
                    word_count=len(content),
                )
            ]

        chapters = self._filter_and_cap(chapters)
        return chapters, len(chapters), {}

    def _parse_txt(self, path: Path) -> Tuple[List[Chapter], int, dict]:
        """解析 TXT 文件"""
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="gbk")

        all_text = [(1, content)]
        chapters = self._parse_chapters_with_pattern(all_text, TXT_CHAPTER_PATTERNS)
        chapters = self._filter_and_cap(chapters)
        return chapters, 1, {}

    def _filter_and_cap(self, chapters: List[Chapter]) -> List[Chapter]:
        """过滤无效章节（封面/版权/前言/目录等）并限制总数"""
        if not chapters:
            return chapters

        filtered = []
        for c in chapters:
            title = (c.title or "").strip()
            if NON_CONTENT_TITLE_RE.match(title):
                continue
            if len(c.content or "") < 200:
                continue
            filtered.append(c)

        if not filtered:
            return chapters[: MAX_CHAPTERS_PER_BOOK]

        if len(filtered) > MAX_CHAPTERS_PER_BOOK:
            filtered.sort(key=lambda c: len(c.content or ""), reverse=True)
            filtered = filtered[:MAX_CHAPTERS_PER_BOOK]
            filtered.sort(key=lambda c: int(c.chapter_number) if c.chapter_number.isdigit() else 9999)

        for i, c in enumerate(filtered):
            c.chapter_number = str(i + 1)

        return filtered


file_parser = FileParser()
